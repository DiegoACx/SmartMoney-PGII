"""
SmartMoney – Fase 6A
Prepara las preguntas REALES del formulario para el etiquetado manual.

Lee ia/fase6/formulario_preguntas_respuestas.xlsx. Las columnas 5 a 9 son las preguntas
abiertas; cada fila es un respondente. Arma una fila por pregunta y marca `descartar`
cuando la respuesta esté vacía o no tenga contenido útil ("si", "nada", "nan", etc.).

Salidas (en ia/fase6/):
  - formulario_preguntas.csv   -> TODAS las preguntas (referencia), con pregunta_num.
  - para_etiquetar_A.csv       -> no descartadas, orden aleatorio (semilla fija), SIN pregunta_num.
  - para_etiquetar_B.csv       -> las mismas, en OTRO orden, para el segundo etiquetador.

Los archivos A y B NO se usan para entrenar: son SOLO para la prueba final (etiquetado humano).

Uso: python preparar_formulario.py
"""
import csv
import os
import random

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(AQUI, "formulario_preguntas_respuestas.xlsx")

COL_PREGUNTAS = [5, 6, 7, 8, 9]          # columnas 1-based con las preguntas
SEMILLA_A, SEMILLA_B = 42, 2025           # órdenes distintos y reproducibles

# Respuestas sin contenido: se descartan aunque tengan texto.
SIN_CONTENIDO = {
    "si", "sí", "no", "nada", "nan", "none", "null", "ninguno", "ninguna", "na", "n/a",
    "-", "--", ".", "..", "x", "xx", "ns", "nose", "no se", "n/a", "0", "null",
}


def limpiar(v):
    """Convierte el valor de la celda a texto limpio ('' si no hay nada)."""
    if v is None:
        return ""
    s = str(v).strip()
    return " ".join(s.split())


def es_descartable(texto):
    """True si la respuesta está vacía o no aporta contenido."""
    if texto == "":
        return True
    if len(texto) < 3:
        return True
    if texto.lower() in SIN_CONTENIDO:
        return True
    return False


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[wb.sheetnames[0]]

    filas = []      # todas las preguntas (id, respondente, pregunta_num, texto, descartar)
    respondente = 0
    contador = 0
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if all(limpiar(fila[c - 1]) == "" for c in COL_PREGUNTAS):
            continue                                  # fila sin ninguna pregunta: se ignora
        respondente += 1
        for pregunta_num, col in enumerate(COL_PREGUNTAS, start=1):
            texto = limpiar(fila[col - 1])
            contador += 1
            filas.append({
                "id": f"P{contador:03d}",
                "respondente": f"R{respondente:03d}",
                "pregunta_num": pregunta_num,
                "texto": texto,
                "descartar": es_descartable(texto),
            })

    # 1) TODAS las preguntas (referencia) con pregunta_num
    ruta_todas = os.path.join(AQUI, "formulario_preguntas.csv")
    with open(ruta_todas, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["id", "respondente", "pregunta_num", "texto", "descartar"])
        w.writeheader()
        for r in filas:
            w.writerow({**r, "descartar": str(r["descartar"]).lower()})

    # 2) Solo las no descartadas, en dos órdenes distintos (A y B)
    utiles = [r for r in filas if not r["descartar"]]
    for nombre, semilla in (("para_etiquetar_A.csv", SEMILLA_A), ("para_etiquetar_B.csv", SEMILLA_B)):
        orden = utiles[:]
        random.Random(semilla).shuffle(orden)
        with open(os.path.join(AQUI, nombre), "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["id", "texto", "intencion"])
            for r in orden:
                w.writerow([r["id"], r["texto"], ""])

    descartadas = len(filas) - len(utiles)
    print(f"Respondentes: {respondente} | preguntas totales: {len(filas)}")
    print(f"Para etiquetar: {len(utiles)} | descartadas: {descartadas}")
    print(f"-> {ruta_todas}")
    print(f"-> para_etiquetar_A.csv (semilla {SEMILLA_A}) y para_etiquetar_B.csv (semilla {SEMILLA_B})")


if __name__ == "__main__":
    main()
